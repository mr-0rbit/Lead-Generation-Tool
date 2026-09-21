#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rynex Security - Automated Lead Generation & Follow-up System  (v2)
===================================================================
For every company it finds, the script works out:
  * which service they probably need (VAPT / SOC / GRC) and WHY (the "hint"),
  * when that hint was posted (job post date / news date, when available),
  * roughly how many employees they have,
  * who the decision-maker is (name, title, email, LinkedIn - only if publicly listed),
then saves it all to an Excel table, data\\leads.xlsx (+ .db; CSV optional), sends an intro email written for THAT
context using your HTML template (email_template.html), and sends timed follow-ups.
It stops automatically on reply / bounce / unsubscribe.

Commands:  init | doctor | find | enrich | preview | send | replies | run | daemon
           export | stats | import | unsubscribe | test-email
"""
from __future__ import annotations

import argparse
import csv
import html as htmllib
import imaplib
import json
import logging
import os
import random
import re
import shutil
import smtplib
import sqlite3
import ssl
import sys
import time
from datetime import datetime, timedelta
from email import message_from_bytes
from email.header import decode_header, make_header
from email.message import EmailMessage
from email.utils import formataddr, formatdate, make_msgid, parseaddr
from pathlib import Path
from urllib import robotparser
from urllib.parse import quote_plus, urljoin, urlparse

try:
    import requests
    from bs4 import BeautifulSoup
except ImportError:
    sys.exit("Missing packages. Run:  pip install -r requirements.txt")

for _s in (sys.stdout, sys.stderr):  # avoid Windows console encoding crashes
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

CONFIG_VERSION = 3
RUNTIME = {"csv": False, "national": set()}  # filled from config at load time
BASE_DIR = Path(os.environ.get("RYNEX_HOME", Path(__file__).resolve().parent))
DATA_DIR = BASE_DIR / "data"
CONFIG_PATH = BASE_DIR / "config.json"
DB_PATH = DATA_DIR / "leads.db"
CSV_PATH = DATA_DIR / "leads.csv"
XLSX_PATH = DATA_DIR / "leads.xlsx"
LOG_PATH = DATA_DIR / "rynex_leadgen.log"

log = logging.getLogger("rynex")

# settings that are reset to the new defaults when an older config.json is upgraded
RESET_ON_UPGRADE = [("templates",), ("services",), ("hooks",), ("industry_pitches",), ("industry_needs",),
                    ("size_lines",), ("compliance",), ("discovery", "industries"), ("discovery", "locations"),
                    ("discovery", "query_templates"), ("signals", "search_jobs")]


_NEEDS = [
    (("bank", "microfinance"), ["vapt", "soc", "grc", "audit"],
     "regulated financial institution and high-value target, expected to run regular penetration tests, 24/7 monitoring and compliance audits",
     "For financial institutions, penetration testing, SOC monitoring and regulatory-aligned GRC support need to work together so findings turn into fixes and audit evidence."),
    (("insurance",), ["vapt", "grc", "audit"],
     "holds large volumes of sensitive customer and claims data under strict regulation",
     "Insurers hold large volumes of sensitive customer and claims data, so regular testing and audit-ready controls matter to regulators and customers alike."),
    (("fintech", "payment", "crypto", "wallet"), ["vapt", "soc", "grc"],
     "handles payments and customer funds through APIs and apps, where one flaw can mean fraud and failed audits (PCI DSS / ISO 27001)",
     "For fintech and payment companies, one exploitable flaw in an API or web app can mean exposed customer data and failed audits."),
    (("health", "hospital", "pharma", "clinic"), ["vapt", "grc", "audit"],
     "holds highly sensitive patient data across applications and integrations",
     "Healthcare organisations hold some of the most sensitive data there is, and their applications and integrations need to be tested the way a real attacker would."),
    (("e-commerce", "ecommerce", "retail", "marketplace"), ["vapt", "soc"],
     "customer accounts, checkout and payment integrations are constant attack targets",
     "For online and retail businesses, checkout flows, customer accounts and payment integrations are prime targets for attackers."),
    (("education", "university", "school"), ["vapt", "grc", "audit"],
     "holds student and staff data across portals and learning platforms",
     "Education providers run many portals and hold student data, which makes regular testing and clear policies important."),
    (("software", "saas", "it services", "managed service"), ["vapt", "audit", "grc"],
     "customers increasingly ask software vendors for proof of security (pen-test reports, ISO 27001 / SOC 2)",
     "For software and IT companies, web apps, APIs and cloud setups are the main attack surface, and customers increasingly ask for proof of security."),
    (("telecom", "internet service"), ["soc", "vapt", "audit"],
     "large, complex attack surface and critical infrastructure that needs continuous monitoring",
     "Telecom environments have a large, complex attack surface that benefits from continuous monitoring and regular testing."),
    (("logistics", "airline", "hotel", "travel"), ["vapt", "soc"],
     "customer booking and operations systems are exposed online and disruption is costly",
     "Operations and booking platforms are exposed online, and an outage or breach is costly, so testing and monitoring pay for themselves."),
    (("oil", "energy", "utility", "manufactur"), ["soc", "audit", "vapt"],
     "business and operational systems are increasingly targeted and downtime is expensive",
     "Energy and industrial companies face growing attacks on both business and operational systems, so monitoring and independent audits are becoming standard."),
    (("real estate", "developer"), ["vapt", "audit"],
     "customer, payment and property data sit in web portals and CRMs",
     "Property businesses keep customer and payment data in portals and CRMs that are worth testing before someone else does."),
    (("law firm", "accounting", "consulting", "legal"), ["grc", "audit", "vapt"],
     "confidential client data and client security questionnaires",
     "Professional-services firms hold confidential client data and are increasingly asked to evidence their security controls."),
    (("media", "government", "non-profit", "ngo"), ["vapt", "soc", "audit"],
     "public-facing sites and donor / citizen data are frequent targets",
     "Public-facing sites and supporter data are frequent targets, and a regular test and audit keeps them safe."),
]


def _build_needs() -> dict:
    out = {}
    for keys, services, reason, pitch in _NEEDS:
        for k in keys:
            out[k] = {"services": services, "reason": reason, "pitch": pitch}
    out["default"] = {
        "services": ["vapt", "audit"],
        "reason": "every internet-facing business benefits from independent security testing and review",
        "pitch": "Whether you are preparing for a compliance audit, launching a new product, or simply want to know how an attacker would see your systems, it helps to find the weak spots first."}
    return out

# --------------------------------------------------------------------------- #
#  DEFAULT CONFIG  (written to config.json by `init`; edit that file)
# --------------------------------------------------------------------------- #
DEFAULT_CONFIG = {
    "config_version": CONFIG_VERSION,
    "company": {
        "name": "Rynex Security",
        "website": "https://rynexsecurity.com",
        "linkedin": "https://www.linkedin.com/company/rynex-security",
        "instagram": "https://www.instagram.com/rynex.sec",
        "sender_name": "YOUR NAME",
        "sender_title": "Business Development",
        "postal_address": "YOUR BUSINESS ADDRESS, CITY, COUNTRY",
        "booking_link": "",
    },
    "smtp": {
        "host": "smtp.gmail.com", "port": 587, "security": "starttls",  # starttls | ssl | none
        "username": "info@rynexsecurity.com",
        "password": "",  # or set env var RYNEX_SMTP_PASSWORD (use an App Password)
        "from_email": "info@rynexsecurity.com", "reply_to": "",
    },
    "imap": {
        "host": "imap.gmail.com", "port": 993, "username": "", "password": "",
        "folder": "INBOX", "lookback_days": 30,
    },
    "discovery": {
        # every industry / place below is picked at random, so all of them get covered over time
        "industries": ["bank", "microfinance bank", "insurance company", "fintech", "payment processor",
                       "cryptocurrency exchange", "e-commerce", "retail chain", "healthcare", "hospital",
                       "pharmaceutical company", "software house", "SaaS", "IT services",
                       "managed service provider", "telecom", "internet service provider", "logistics",
                       "airline", "hotel group", "education technology", "university", "oil and gas",
                       "energy utility", "manufacturing", "real estate developer", "law firm",
                       "accounting firm", "consulting firm", "media company", "non-profit organization"],
        "national_locations": ["Pakistan", "Karachi", "Lahore", "Islamabad", "Rawalpindi", "Faisalabad",
                               "Multan", "Peshawar", "Quetta", "Sialkot"],
        "international_locations": ["United Arab Emirates", "Dubai", "Saudi Arabia", "Qatar", "Kuwait", "Bahrain",
                                    "Oman", "United Kingdom", "London", "United States", "Canada", "Australia",
                                    "Germany", "Netherlands", "Singapore", "Malaysia", "Turkey", "Egypt",
                                    "Nigeria", "Kenya", "South Africa", "India", "Bangladesh", "Sri Lanka",
                                    "Indonesia", "Philippines", ""],  # "" = worldwide, no place in the search
        "national_ratio": 0.4,  # share of searches that target Pakistan; the rest are international
        "query_templates": ["{industry} company in {location}", "{industry} {location} contact us email",
                            "{industry} {location} IT security"],
        "min_employees": 0, "max_employees": 0,  # 0 = no size limit
        "priority_industries": ["fintech", "bank", "health", "e-commerce", "software", "payment", "telecom"],
        "max_queries_per_run": 8, "results_per_query": 15, "max_new_leads_per_run": 40,
        "requery_after_days": 30, "pause_discovery_if_backlog_over": 150,
        "google_places_api_key": "",  # optional, gives better results
        "request_delay_seconds": [1.5, 3.5],
        "user_agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) RynexLeadBot/1.0 (+https://rynexsecurity.com)",
        "extra_blocked_domains": [],
    },
    "signals": {  # how the "need hint", employee count etc. are found
        "enabled": True,
        "search_news": True,        # dated news: incidents, compliance, launches, funding
        "use_hiring_hints": False,  # True = also treat security job posts as a buying signal
        "search_employees": True,   # employee-count lookup via search snippets
        "only_for_leads_with_email": True,
        "hint_max_age_days": 365,   # ignore dated hints older than this
        "search_delay_seconds": [2, 4],
        "cite_incident_news": False,  # False = never mention a breach in the email itself (safer)
    },
    "sending": {
        "daily_limit": 25, "delay_between_emails_seconds": [60, 150], "min_lead_score": 40,
        "followup_days": [3, 7, 14],
        "send_window": {"days": [0, 1, 2, 3, 4], "start_hour": 10, "end_hour": 16},
    },
    "automation": {"cycle_hours": 6, "notify_email": ""},
    "output": {"excel": True, "csv": False},  # Excel table is the main result; CSV only if you turn it on
    "email_template_file": "email_template.html",
    "compliance": {
        "footer": (
            "You are receiving this one-to-one business email because {company} is publicly listed as a "
            "business at {domain}. If you would rather not hear from us, just reply \"unsubscribe\" and we "
            "will remove you immediately.\n{sender_company} | {postal_address}"
        )
    },
    # ---- what we sell, per service -------------------------------------------------------------
    "services": {
        "vapt": {
            "label": "VAPT", "phrase": "penetration testing",
            "subject": "Penetration testing for {company}",
            "offer": "Rynex Security runs real-world attack simulations (VAPT) against web apps, APIs, networks and cloud environments, so weaknesses are found and fixed before an attacker finds them.",
            "deliverable": "You get a clear, prioritised report your developers can act on straight away, with evidence for every finding and retesting to confirm the fixes.",
            "checklist": "exposed admin panels and forgotten subdomains, weak authentication and broken access control, misconfigured cloud storage, and unpatched internet-facing software",
        },
        "soc": {
            "label": "SOC", "phrase": "24/7 security monitoring",
            "subject": "24/7 security monitoring for {company}",
            "offer": "Our Security Operations Center provides continuous monitoring, threat detection and rapid incident response, so suspicious activity is spotted and contained in minutes rather than days.",
            "deliverable": "That means alerts triaged by certified analysts, clear escalation, and proactive threat hunting, without the cost of building a full in-house SOC.",
            "checklist": "which log sources are actually monitored, how quickly an alert reaches a human, whether there is a tested incident-response playbook, and whether endpoints and cloud accounts are covered",
        },
        "grc": {
            "label": "GRC", "phrase": "compliance readiness",
            "subject": "Compliance readiness for {company}",
            "offer": "Our GRC team helps organisations align with standards such as ISO 27001, PCI DSS and SOC 2, and manage risk in a way that produces real audit evidence, not just paperwork.",
            "deliverable": "We start with a gap assessment, then give you a practical roadmap covering policies, controls and evidence collection.",
            "checklist": "scope and asset inventory, risk assessment, access-control and logging evidence, vendor risk, and incident-response documentation",
        },
        "audit": {
            "label": "Security Audit", "phrase": "an independent security audit",
            "subject": "Independent security audit for {company}",
            "offer": "Our security audits give you an independent review of your configurations, architecture, access controls and policies against recognised standards such as ISO 27001, PCI DSS and NIST, so gaps are documented before a customer, regulator or attacker finds them.",
            "deliverable": "You receive a prioritised findings report with risk ratings, evidence and a practical remediation plan.",
            "checklist": "firewall and cloud configuration reviews, privileged-access and identity controls, backup and recovery readiness, logging coverage, and gaps between written policy and actual practice",
        },
        "general": {
            "label": "General security", "phrase": "security testing and monitoring",
            "subject": "Quick question about security at {company}",
            "offer": "Rynex Security helps organisations find and fix real security weaknesses before attackers do, through penetration testing (VAPT), 24/7 SOC monitoring, GRC / compliance support and security audits.",
            "deliverable": "You get certified professionals with an offensive-security mindset, a sharp focus on the risks that matter, and clear, actionable reports.",
            "checklist": "exposed admin panels and forgotten subdomains, weak authentication, misconfigured cloud storage, and unpatched internet-facing software",
        },
    },
    # ---- opening line, chosen by the type of hint we found -------------------------------------------
    "hooks": {
        "hiring": "I noticed {company} is hiring for the role of {hint_detail}. That usually means security is becoming a bigger priority, and we often help teams cover the gap while they recruit.",
        "incident": "I came across public reports of a security incident involving {company}, and I'm sorry you've had to deal with that. When teams are rebuilding trust, an independent review can help confirm nothing was missed.",
        "compliance": "I saw that {company} has been in the news around compliance and certification ({hint_detail}). Getting audit-ready is much easier with an outside view early on.",
        "launch": "I saw the recent news about {company} ({hint_detail}). Congratulations! New products and growth are exactly when an independent security review pays off most.",
        "website_compliance": "{company} references {hint_detail} on its website, so I imagine security assurance matters to your customers.",
        "industry": "{industry_pitch}",
    },
    "size_lines": {
        "small": "We work with lean teams, so engagements are scoped to fit your size and budget.",
        "medium": "For a team of your size, we can start with a focused, time-boxed engagement and expand from there.",
        "large": "For larger environments, we plan around your existing tools and teams so testing and monitoring fit your current processes.",
        "unknown": "",
    },
    "industry_needs": _build_needs(),  # per industry: which services they need, why, and the email pitch
    "templates": {
        "initial": {
            "subject": "{subject}",
            "body": (
                "{greeting}\n\n{hook}\n\n{service_offer} {size_line}\n\n{service_deliverable}\n\n"
                "{forward_line}\n\n"
                "Would you be open to a short 15-minute call to see whether we can help? {booking_line}\n\n"
                "Best regards,\n{sender_name}\n{sender_title}, Rynex Security\n{website}"
            ),
        },
        "followups": [
            {"subject": "", "body": (
                "{greeting}\n\nJust following up on my note from a few days ago about {service_phrase}. "
                "{service_deliverable}\n\nIf this is on your roadmap, I'd be glad to share how we would approach it "
                "for {company}. {booking_line}\n\nBest regards,\n{sender_name}\n{sender_title}, Rynex Security\n{website}")},
            {"subject": "", "body": (
                "{greeting}\n\nCommon starting points when we look at {service_phrase}: {service_checklist}.\n\n"
                "Happy to do a quick, no-obligation scoping chat if any of that sounds relevant. {booking_line}\n\n"
                "Best regards,\n{sender_name}\n{sender_title}, Rynex Security\n{website}")},
            {"subject": "", "body": (
                "{greeting}\n\nI don't want to clutter your inbox, so this will be my last message. If the timing "
                "isn't right, no problem at all. You can learn more at {website} or reach us any time at "
                "{from_email}.\n\nWishing {company} all the best.\n\n"
                "Best regards,\n{sender_name}\n{sender_title}, Rynex Security")},
        ],
    },
}

# Your HTML design. Written to email_template.html so you can edit it freely.
# Placeholders: {{PREHEADER}} {{BODY}} {{FOOTER}} {{CONTACT_EMAIL}} {{WEBSITE}}
DEFAULT_EMAIL_HTML = """<div>
<div style="display:none;max-height:0px;opacity:0;overflow:hidden;">{{PREHEADER}}</div>
<table style="padding:24px 12px;background-color:rgb(244,246,248);" border="0" cellpadding="0" cellspacing="0" width="100%">
<tbody><tr><td align="center">
<table style="max-width:640px;border-radius:16px;overflow:hidden;box-shadow:0px 10px 30px rgba(0,0,0,0.12);background-color:rgb(255,255,255);" border="0" cellpadding="0" cellspacing="0" width="100%">
<tbody>
<tr><td style="padding:24px 32px;text-align:center;background-color:rgb(0,0,0);">
<img style="display:block;margin:0 auto 12px auto;border-radius:12px" height="72" width="72" alt="Rynex Security Logo" src="https://ik.imagekit.io/t4itchmhb/logo.png">
<div style="font-size:28px;font-weight:bold;letter-spacing:0.4px;color:rgb(255,255,255);">Rynex Security</div>
<div style="font-size:15px;font-weight:bold;margin-top:6px;color:rgb(0,194,255);">Detect &bull; Exploit &bull; Secure</div>
</td></tr>
<tr><td style="padding:32px;font-family:'Segoe UI',Arial,sans-serif;">
{{BODY}}
<p style="margin:0px;line-height:1.7;"><span style="color:rgb(55,65,81);font-size:16px;">If you have any questions, please contact us at <a target="_blank" style="text-decoration:none;color:rgb(37,99,235);" href="mailto:{{CONTACT_EMAIL}}">{{CONTACT_EMAIL}}</a>.</span></p>
</td></tr>
<tr><td style="padding:24px 32px;text-align:center;background-color:rgb(0,0,0);">
<div style="font-size:20px;font-weight:bold;margin-bottom:6px;color:rgb(255,255,255);">Rynex Security</div>
<div style="font-size:14px;margin-bottom:8px;color:rgb(203,213,225);">Offensive Security &bull; Penetration Testing &bull; Cloud Security</div>
<div><a style="display:inline-block;text-decoration:none;font-size:14px;color:rgb(0,194,255);" target="_blank" href="{{WEBSITE}}">{{WEBSITE}}</a></div>
<div style="margin-top:16px;font-size:13px;line-height:1.7;color:rgb(156,163,175);">{{FOOTER}}</div>
</td></tr>
</tbody></table>
</td></tr></tbody></table>
</div>"""

# --------------------------------------------------------------------------- #
#  CONSTANTS
# --------------------------------------------------------------------------- #
BLOCKED_DOMAINS = {
    "linkedin.com", "facebook.com", "instagram.com", "twitter.com", "x.com", "youtube.com",
    "wikipedia.org", "reddit.com", "quora.com", "medium.com", "pinterest.com", "tiktok.com",
    "clutch.co", "goodfirms.co", "themanifest.com", "crunchbase.com", "glassdoor.com",
    "indeed.com", "yelp.com", "tripadvisor.com", "trustpilot.com", "g2.com", "capterra.com",
    "zoominfo.com", "rocketreach.co", "apollo.io", "dnb.com", "mapquest.com", "amazon.com",
    "forbes.com", "techcrunch.com", "businessinsider.com", "yellowpages.com", "justdial.com",
    "google.com", "bing.com", "github.com", "gitlab.com", "olx.com.pk", "rozee.pk",
    "wixsite.com", "wordpress.com", "blogspot.com", "sentry.io", "godaddy.com", "sbp.org.pk",
}
BLOCKED_SUFFIXES = (".gov", ".mil", ".gov.pk", ".gov.uk", ".gov.ae", ".gov.sa", ".edu")
FREE_MAIL = {"gmail.com", "yahoo.com", "outlook.com", "hotmail.com", "live.com", "icloud.com"}
JUNK_LOCALS = {
    "noreply", "no-reply", "donotreply", "do-not-reply", "mailer-daemon", "postmaster",
    "example", "test", "yourname", "name", "user", "email", "username", "your", "abuse",
    "privacy", "unsubscribe", "root", "john.doe", "jane.doe",
}
JUNK_EMAIL_DOMAINS = {"example.com", "sentry.io", "wixpress.com", "sentry-next.wixpress.com",
                      "domain.com", "email.com", "yourdomain.com", "godaddy.com"}
BAD_EXT = (".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp", ".css", ".js", ".ico", ".woff", ".woff2")
PREFERRED_LOCALS = ["ceo", "founder", "cto", "ciso", "security", "it", "info", "contact",
                    "hello", "sales", "business", "admin", "office", "support"]
GENERIC_LOCALS = ["info", "contact", "hello", "sales", "business", "office", "admin", "support"]
CONTACT_HINTS = ("contact", "about", "team", "reach", "connect", "company", "get-in-touch", "impressum",
                 "leadership", "management", "our-people", "people")
CAREER_HINTS = ("career", "jobs", "join-us", "join our", "hiring", "vacanc", "work-with-us", "opportunit")
NEED_KEYWORDS = ("payment", "customer data", "api", "cloud", "checkout", "login", "patient",
                 "saas", "portal", "mobile app", "e-commerce", "wallet", "online banking")
FRAMEWORKS = [("iso 27001", "ISO 27001"), ("iso27001", "ISO 27001"), ("pci dss", "PCI DSS"), ("pci-dss", "PCI DSS"),
              ("soc 2", "SOC 2"), ("gdpr", "GDPR"), ("hipaa", "HIPAA")]
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9\-]+(?:\.[A-Za-z0-9\-]+)*\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:\+|00)\d{1,3}[\s\-.]?\(?\d{2,4}\)?[\s\-.]?\d{3,4}[\s\-.]?\d{3,4}|\b0\d{2,3}[\s\-]?\d{7,8}\b")
UNSUB_PHRASES = ("unsubscribe", "remove me", "remove us", "stop emailing", "stop sending",
                 "do not contact", "don't contact", "dont contact", "opt out", "opt-out", "take me off")
DECLINE_PHRASES = ("not interested", "no thanks", "no thank you", "not looking", "please don't follow")

SEC_JOB_RE = re.compile(
    r"\b(?:pen(?:etration)?[\s\-]?test\w*|red[\s\-]?team\w*|ciso|devsecops|appsec|head of (?:information |cyber ?)?security|"
    r"(?:cyber[\s\-]?security|information security|infosec|security|soc|threat|vulnerability|incident response|"
    r"compliance|grc|network security|cloud security|application security)\s+"
    r"(?:engineer|analyst|specialist|officer|manager|lead|architect|consultant|researcher|tester|administrator|"
    r"director|head|hunter|expert|associate)s?)\b", re.I)
JOB_GRC_RE = re.compile(r"compliance|grc|governance|audit|iso ?27001|isms|ciso|security officer|risk", re.I)
JOB_VAPT_RE = re.compile(r"pen\w* ?test|red[\s\-]?team|appsec|application security|vulnerab|devsecops|security engineer|researcher|tester", re.I)
INCIDENT_RE = re.compile(r"breach|cyber ?attack|ransomware|hacked|data leak|leaked|security incident|compromised|malware|phishing attack", re.I)
COMPLIANCE_RE = re.compile(r"iso ?27001|pci[\s\-]?dss|soc ?2|gdpr|hipaa|compliance|certif(?:ied|ication)|audit", re.I)
LAUNCH_RE = re.compile(r"launch|unveil|introduces|new app|new platform|raises|funding|series [abc]\b|acquir|expands|expansion|partnership|go[\s\-]live", re.I)
TITLE_RE = re.compile(
    r"\b(ceo|chief executive|founder|co-?founder|managing director|owner|president|cto|chief technology|ciso|"
    r"chief information security|head of (?:it|security|technology|engineering|infosec|information security)|"
    r"it (?:manager|director|head)|director of (?:it|technology|information technology|security)|"
    r"(?:information|cyber)\s?security (?:manager|officer|head|lead|director)|security (?:manager|head|lead|director|officer)|"
    r"compliance (?:manager|officer|head)|coo|chief operating|vp (?:of )?engineering|technical director|"
    r"technology officer|chief information officer|cio)\b", re.I)
NAME_RE = re.compile(r"^[A-Z][a-z'\u2019\-]+(?:\s+[A-Z][a-z'\u2019\-\.]*){1,3}$")
STOP_TOKENS = {"our", "team", "meet", "the", "contact", "us", "about", "read", "more", "learn", "view", "services",
               "company", "leadership", "management", "board", "director", "directors", "message", "from", "with",
               "and", "get", "started", "home", "careers", "blog", "news", "privacy", "policy", "terms", "cyber",
               "security", "technology", "technologies", "solutions", "chief", "officer", "executive", "managing",
               "head", "manager", "founder", "cofounder", "owner", "president", "it", "information", "engineering",
               "operating", "technical", "compliance", "senior", "lead", "experts", "expert"}
BUYER_ORDER = {
    "default": [r"ciso|chief information security|head of (?:information |cyber )?security|(?:information |cyber )?security (?:manager|head|director|lead|officer)|infosec",
                r"cto|chief technology|technology officer|technical director|head of (?:technology|engineering|it)|vp|it (?:manager|director|head)|director of (?:it|technology|information technology)|chief information officer|cio",
                r"compliance", r"ceo|chief executive|founder|managing director|owner|president", r"coo|chief operating"],
    "grc": [r"compliance", r"ciso|chief information security|(?:information |cyber )?security (?:manager|head|director|lead|officer)|infosec",
            r"ceo|chief executive|founder|managing director|owner|president", r"cto|chief technology|it (?:manager|director|head)|head of it"],
}
NON_SALES_TOKENS = {"support", "care", "hr", "job", "jobs", "resume", "cv", "billing", "invoice", "invoices", "payroll",
                    "press", "media", "newsletter", "privacy", "dpo", "legal", "fraud", "abuse", "report", "reports",
                    "whistleblower", "helpdesk", "customercare", "complaint", "complaints"}
NON_SALES_PARTS = ("customer", "complain", "helpdesk", "recruit", "career", "remittance", "feedback", "grievance")
LEGAL_SUFFIX = re.compile(r"\b(pvt|private|ltd|limited|llc|inc|corp|corporation|co|company|smc|plc|gmbh|fze|fzco)\b\.?", re.I)

SCHEMA = """
CREATE TABLE IF NOT EXISTS leads (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  domain TEXT UNIQUE, company TEXT, website TEXT, industry TEXT, location TEXT,
  email TEXT, all_emails TEXT, phone TEXT, linkedin TEXT, address TEXT, description TEXT,
  source TEXT, score INTEGER DEFAULT 0, status TEXT DEFAULT 'new', step INTEGER DEFAULT 0,
  fail_count INTEGER DEFAULT 0, last_sent_at TEXT, next_send_at TEXT,
  first_message_id TEXT, last_message_id TEXT, replied_at TEXT, notes TEXT, created_at TEXT
);
CREATE TABLE IF NOT EXISTS emails_sent (
  id INTEGER PRIMARY KEY AUTOINCREMENT, lead_id INTEGER, step INTEGER, to_email TEXT,
  subject TEXT, message_id TEXT, sent_at TEXT
);
CREATE TABLE IF NOT EXISTS suppression (email TEXT PRIMARY KEY, reason TEXT, added_at TEXT);
CREATE TABLE IF NOT EXISTS seen_queries (query TEXT PRIMARY KEY, run_at TEXT);
"""
# columns added in v2 (auto-migrated into an existing database)
NEW_COLS = {
    "company_email": "TEXT", "need_service": "TEXT", "need_services_all": "TEXT", "hint_type": "TEXT",
    "hint_detail": "TEXT", "hint_text": "TEXT", "hint_date": "TEXT", "hint_url": "TEXT", "hint_found_on": "TEXT",
    "other_hints": "TEXT", "signals_json": "TEXT", "employees": "TEXT", "employees_est": "INTEGER",
    "employees_source": "TEXT", "contact_name": "TEXT", "contact_title": "TEXT", "contact_email": "TEXT",
    "contact_linkedin": "TEXT", "buyer_search_url": "TEXT", "other_contacts": "TEXT", "market": "TEXT",
}

EXPORT_COLUMNS = [
    ("Company", "company"), ("Website", "website"), ("Market", "market"), ("Location", "location"),
    ("Industry", "industry"), ("Service Needed", "need_services_all"),
    ("Need Hint (why they need it)", "hint_text"), ("Hint Type", "hint_type"),
    ("Hint Posted Date", "hint_date"), ("Hint Found On", "hint_found_on"), ("Hint Source URL", "hint_url"),
    ("Other Hints", "other_hints"), ("Employees (current)", "employees"), ("Employees Source", "employees_source"),
    ("Decision-Maker", "contact_name"), ("Decision-Maker Title", "contact_title"),
    ("Decision-Maker Email", "contact_email"), ("Decision-Maker LinkedIn", "contact_linkedin"),
    ("Other Leaders Found", "other_contacts"), ("Find Decision-Maker (LinkedIn search)", "buyer_search_url"),
    ("Company Email", "company_email"), ("Company Phone", "phone"), ("Company LinkedIn", "linkedin"),
    ("Send-To Email", "email"), ("Other Emails Found", "all_emails"), ("Lead Score", "score"),
    ("Status", "status"), ("Emails Sent", "step"), ("Last Emailed", "last_sent_at"),
    ("Next Follow-up", "next_send_at"), ("Replied At", "replied_at"), ("Source", "source"), ("Notes", "notes"),
]


# --------------------------------------------------------------------------- #
#  HELPERS
# --------------------------------------------------------------------------- #
def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def deep_merge(base: dict, over: dict) -> dict:
    out = dict(base)
    for k, v in over.items():
        out[k] = deep_merge(base[k], v) if isinstance(v, dict) and isinstance(base.get(k), dict) else v
    return out


def _pop_path(d: dict, path: tuple):
    for k in path[:-1]:
        d = d.get(k)
        if not isinstance(d, dict):
            return
    d.pop(path[-1], None)


def load_config() -> dict:
    if not CONFIG_PATH.exists():
        sys.exit(f"config.json not found. Run:  python {Path(__file__).name} init")
    with open(CONFIG_PATH, encoding="utf-8") as f:
        raw = json.load(f)
    old = raw.get("config_version", 1)
    if old < CONFIG_VERSION:
        backup = CONFIG_PATH.with_name(f"config.backup_v{old}.json")
        shutil.copy(CONFIG_PATH, backup)
        for path in RESET_ON_UPGRADE:  # old wording / search areas are replaced by the new defaults
            _pop_path(raw, path)
        raw["config_version"] = CONFIG_VERSION
        CONFIG_PATH.write_text(json.dumps(deep_merge(DEFAULT_CONFIG, raw), indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"[upgrade] config.json upgraded to v{CONFIG_VERSION}: your email/SMTP settings are kept; search areas are "
              f"now national + international, and email wording/services were refreshed. Old file: {backup.name}")
    cfg = deep_merge(DEFAULT_CONFIG, raw)
    cfg["smtp"]["password"] = os.environ.get("RYNEX_SMTP_PASSWORD") or cfg["smtp"]["password"]
    im = cfg["imap"]
    im["username"] = im["username"] or cfg["smtp"]["username"]
    im["password"] = os.environ.get("RYNEX_IMAP_PASSWORD") or im["password"] or cfg["smtp"]["password"]
    RUNTIME["csv"] = bool(cfg["output"]["csv"])
    RUNTIME["national"] = {x.lower() for x in cfg["discovery"]["national_locations"]}
    return cfg


def setup_logging():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    fmt = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s", "%Y-%m-%d %H:%M:%S")
    log.setLevel(logging.INFO)
    if not log.handlers:
        for h in (logging.StreamHandler(sys.stdout), logging.FileHandler(LOG_PATH, encoding="utf-8")):
            h.setFormatter(fmt)
            log.addHandler(h)


def sleep_range(rng):
    time.sleep(random.uniform(rng[0], rng[1]))


def root_domain(url: str) -> str:
    try:
        host = urlparse(url if "//" in url else "//" + url).hostname or ""
    except ValueError:
        return ""
    return host.lower().removeprefix("www.")


def is_blocked(domain: str, cfg: dict) -> bool:
    extra = set(cfg["discovery"].get("extra_blocked_domains", []))
    if domain in BLOCKED_DOMAINS or domain in extra:
        return True
    if any(domain.endswith("." + b) for b in BLOCKED_DOMAINS | extra):
        return True
    return domain.endswith(BLOCKED_SUFFIXES) or ("." not in domain and domain != "localhost")


def decode_hdr(v) -> str:
    try:
        return str(make_header(decode_header(v or "")))
    except Exception:
        return v or ""


def parse_date(s) -> str:
    """Best-effort -> 'YYYY-MM-DD' or ''."""
    if not s:
        return ""
    s = str(s).strip()
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        m = re.search(r"\d{4}-\d{2}-\d{2}", s)
        return m.group(0) if m else ""


def too_old(date_str: str, max_days: int) -> bool:
    if not date_str:
        return False
    try:
        return (datetime.now().date() - datetime.fromisoformat(date_str).date()).days > max_days
    except ValueError:
        return False


class SafeDict(dict):
    def __missing__(self, key):
        return "{" + key + "}"


# --------------------------------------------------------------------------- #
#  DATABASE
# --------------------------------------------------------------------------- #
class DB:
    def __init__(self, path: Path = DB_PATH):
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(path)
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(SCHEMA)
        have = {r["name"] for r in self.conn.execute("PRAGMA table_info(leads)")}
        for col, typ in NEW_COLS.items():
            if col not in have:
                self.conn.execute(f"ALTER TABLE leads ADD COLUMN {col} {typ}")
        self.conn.commit()

    def all(self, sql, args=()):
        return self.conn.execute(sql, args).fetchall()

    def one(self, sql, args=()):
        return self.conn.execute(sql, args).fetchone()

    def run(self, sql, args=()):
        cur = self.conn.execute(sql, args)
        self.conn.commit()
        return cur

    def update_lead(self, lead_id: int, **fields):
        cols = ", ".join(f"{k}=?" for k in fields)
        self.run(f"UPDATE leads SET {cols} WHERE id=?", (*fields.values(), lead_id))

    def add_lead(self, lead: dict) -> bool:
        lead.setdefault("created_at", now_iso())
        cols = ",".join(lead)
        marks = ",".join("?" * len(lead))
        cur = self.run(f"INSERT OR IGNORE INTO leads ({cols}) VALUES ({marks})", tuple(lead.values()))
        return cur.rowcount > 0

    def lead_exists(self, domain: str) -> bool:
        return self.one("SELECT 1 FROM leads WHERE domain=?", (domain,)) is not None

    def suppress(self, email: str, reason: str):
        self.run("INSERT OR IGNORE INTO suppression VALUES (?,?,?)", (email.lower(), reason, now_iso()))

    def is_suppressed(self, email: str) -> bool:
        return self.one("SELECT 1 FROM suppression WHERE email=?", ((email or "").lower(),)) is not None


# --------------------------------------------------------------------------- #
#  SEARCH (web search + news search via the `ddgs` package)
# --------------------------------------------------------------------------- #
SEARCH_STATS = {"queries": 0, "results": 0, "errors": 0, "last_error": ""}


def ddg_search(kind: str, query: str, max_results: int, timelimit: str | None = None) -> list[dict]:
    """kind = 'text' or 'news'. Never raises; failures are counted in SEARCH_STATS."""
    SEARCH_STATS["queries"] += 1
    for attempt in (1, 2):
        try:
            try:
                from ddgs import DDGS
            except ImportError:
                from duckduckgo_search import DDGS
            kw = {"max_results": max_results}
            if timelimit:
                kw["timelimit"] = timelimit
            d = DDGS()
            res = (d.news if kind == "news" else d.text)(query, **kw) or []
            SEARCH_STATS["results"] += len(res)
            return res
        except ImportError:
            SEARCH_STATS["errors"] += 1
            SEARCH_STATS["last_error"] = "search package not installed (pip install -U ddgs)"
            return []
        except Exception as e:
            msg = str(e)
            if "no results" in msg.lower():
                return []
            SEARCH_STATS["errors"] += 1
            SEARCH_STATS["last_error"] = msg[:200]
            if attempt == 1 and re.search(r"rate|429|202|limit|timeout", msg, re.I):
                log.warning("Search rate-limited/timeout, waiting 20s then retrying once...")
                time.sleep(20)
                continue
            log.warning("Search failed for '%s': %s", query, msg[:150])
            return []
    return []


def search_ddg(query: str, n: int) -> list[dict]:
    res = ddg_search("text", query, n)
    return [{"url": r.get("href") or r.get("url"), "name": "", "phone": "", "address": "", "source": "web-search"}
            for r in res if r.get("href") or r.get("url")]


def search_places(query: str, key: str, n: int) -> list[dict]:
    try:
        r = requests.post(
            "https://places.googleapis.com/v1/places:searchText",
            json={"textQuery": query, "pageSize": min(n, 20)},
            headers={"X-Goog-Api-Key": key,
                     "X-Goog-FieldMask": "places.displayName,places.websiteUri,places.nationalPhoneNumber,"
                                         "places.internationalPhoneNumber,places.formattedAddress"},
            timeout=20)
        r.raise_for_status()
    except requests.RequestException as e:
        log.warning("Google Places failed: %s", e)
        return []
    out = []
    for p in r.json().get("places", []):
        if p.get("websiteUri"):
            out.append({"url": p["websiteUri"], "name": (p.get("displayName") or {}).get("text", ""),
                        "phone": p.get("internationalPhoneNumber") or p.get("nationalPhoneNumber") or "",
                        "address": p.get("formattedAddress", ""), "source": "google-places"})
    return out


# --------------------------------------------------------------------------- #
#  PAGE PARSING: emails, employees, jobs, decision-makers
# --------------------------------------------------------------------------- #
def decode_cf_email(hexstr: str) -> str:
    try:
        key = int(hexstr[:2], 16)
        return "".join(chr(int(hexstr[i:i + 2], 16) ^ key) for i in range(2, len(hexstr), 2))
    except ValueError:
        return ""


def rank_email(e: str, domain: str):
    local, dom = e.split("@", 1)
    same = dom == domain or dom.endswith("." + domain)
    try:
        r = PREFERRED_LOCALS.index(local)
    except ValueError:
        r = len(PREFERRED_LOCALS) + 1  # unknown person: role unknown, so role addresses (info@, security@...) come first
    return (0 if same else 1, r)


def is_service_mailbox(local: str) -> bool:
    """customer-care / complaints / HR / support inboxes: useless for a sales pitch."""
    return bool(set(re.split(r"[._\-\d]+", local)) & NON_SALES_TOKENS) or any(p in local for p in NON_SALES_PARTS)


def clean_emails(raw: list[str], domain: str, service_only: bool = False) -> list[str]:
    """service_only=False -> usable business addresses; True -> the customer-service ones we skipped."""
    seen, out = set(), []
    for e in raw:
        e = e.strip().strip(".,;:<>()[]\"'").lower()
        if not EMAIL_RE.fullmatch(e) or e.endswith(BAD_EXT) or e in seen:
            continue
        local, dom = e.split("@", 1)
        if local in JUNK_LOCALS or dom in JUNK_EMAIL_DOMAINS or re.fullmatch(r"[0-9a-f]{20,}", local):
            continue
        same = dom == domain or dom.endswith("." + domain)
        if not same and dom not in FREE_MAIL:
            continue  # third-party addresses in footers (agencies, plugins...) are not the company
        if is_service_mailbox(local) != service_only:
            continue
        seen.add(e)
        out.append(e)
    out.sort(key=lambda e: rank_email(e, domain))
    return out


def jsonld_items(soup) -> list[dict]:
    items: list[dict] = []

    def walk(o):
        if isinstance(o, list):
            for x in o:
                walk(x)
        elif isinstance(o, dict):
            items.append(o)
            for v in o.values():
                if isinstance(v, (list, dict)):
                    walk(v)

    for sc in soup.find_all("script", type="application/ld+json"):
        try:
            walk(json.loads(sc.string or sc.get_text() or ""))
        except (ValueError, TypeError):
            pass
    return items


def _to_int(s) -> int:
    try:
        return int(str(s).replace(",", "").strip())
    except ValueError:
        return 0


def parse_employees(text: str) -> tuple[str, int]:
    """Find a headcount statement. Returns (display, estimate) or ('', 0)."""
    pats = [
        (re.compile(r"(\d[\d,]*)\s*[-\u2013]\s*(\d[\d,]*)\s*\+?\s*(?:employees|team members|staff)", re.I), "range"),
        (re.compile(r"(\d[\d,]*)\s*(\+)?\s*(?:full[- ]time\s+)?(?:employees|team members|staff members|staff|professionals|engineers|experts|specialists|consultants)\b", re.I), "single"),
        (re.compile(r"team of (?:over |more than |about |around )?(\d[\d,]*)(\+)?", re.I), "single"),
        (re.compile(r"(?:over|more than|about|around|approximately)\s+(\d[\d,]*)()\s+(?:employees|people|professionals|team members)", re.I), "single"),
    ]
    for rx, kind in pats:
        for m in rx.finditer(text):
            if kind == "range":
                lo, hi = _to_int(m.group(1)), _to_int(m.group(2))
                if 1 <= lo < hi <= 1_000_000:
                    return f"{m.group(1)}-{m.group(2)}", (lo + hi) // 2
            else:
                n = _to_int(m.group(1))
                if 2 <= n <= 1_000_000:
                    return f"{m.group(1)}{'+' if m.group(2) else ''}", n
    return "", 0


def extract_jobs(soup, page_url: str, text_ok: bool) -> list[dict]:
    jobs, seen = [], set()
    for it in jsonld_items(soup):
        t = it.get("@type")
        if "JobPosting" in (t if isinstance(t, list) else [t]) and it.get("title"):
            title = str(it["title"]).strip()
            if SEC_JOB_RE.search(title) and title.lower() not in seen:
                seen.add(title.lower())
                jobs.append({"title": title[:90], "date": parse_date(it.get("datePosted")),
                             "url": it.get("url") or page_url})
    if text_ok:
        for el in soup.find_all(["h1", "h2", "h3", "h4", "a", "li", "p"]):
            txt = el.get_text(" ", strip=True)
            if 6 <= len(txt) <= 90 and SEC_JOB_RE.search(txt) and txt.lower() not in seen:
                seen.add(txt.lower())
                jobs.append({"title": txt, "date": "", "url": page_url})
    return jobs


def clean_name(s: str) -> str:
    s = re.sub(r"^(dr|mr|ms|mrs|eng|engr|prof)\.?\s+", "", (s or "").strip(), flags=re.I)
    if s.isupper():
        s = s.title()
    if not (5 <= len(s) <= 40) or not NAME_RE.match(s):
        return ""
    if any(t.lower().strip(".") in STOP_TOKENS for t in s.split()):
        return ""
    return s


def extract_people(html: str) -> list[dict]:
    """Names + job titles of leaders listed on a page (only what the company publishes)."""
    soup = BeautifulSoup(html, "html.parser")
    links = []  # (url, [text of each ancestor level, smallest first])
    for a in soup.find_all("a", href=re.compile(r"linkedin\.com/in/", re.I)):
        ctxs, node = [], a
        for _ in range(4):
            node = node.parent
            if node is None:
                break
            ctxs.append(node.get_text(" ", strip=True).lower()[:600])
        links.append((a["href"].split("?")[0], ctxs))
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    strings = [s.strip() for s in soup.stripped_strings]
    people, seen = [], set()
    for i, s in enumerate(strings):
        if len(s) > 70 or not TITLE_RE.search(s):
            continue
        name = title = ""
        parts = [p for p in re.split(r"\s*(?:,|\||\u2022)\s*|\s+[\-\u2013\u2014/]\s+", s) if p]
        if len(parts) >= 2:
            for j, part in enumerate(parts):
                nm = clean_name(part)
                rest = " ".join(p for k, p in enumerate(parts) if k != j)
                if nm and TITLE_RE.search(rest):
                    name, title = nm, rest.strip()
                    break
        if not name and len(s) < 60:
            title = s
            for cand in (strings[i - 1] if i > 0 else "", strings[i + 1] if i + 1 < len(strings) else ""):
                nm = clean_name(cand)
                if nm:
                    name = nm
                    break
        if name and title and name.lower() not in seen:
            seen.add(name.lower())
            people.append({"name": name, "title": title[:70], "linkedin": ""})
    # attach LinkedIn profiles: by name in the URL slug, else by the closest block that mentions exactly one person
    def toks(nm):
        t = [x.lower() for x in nm.split()]
        return t[0], t[-1]

    for href, ctxs in links:
        slug = href.lower()
        owner = next((p for p in people if all(x in slug for x in toks(p["name"]))), None)
        if owner is None:
            for ctx in ctxs:
                hit = [p for p in people if all(x in ctx for x in toks(p["name"]))]
                if len(hit) == 1:
                    owner = hit[0]
                    break
                if len(hit) > 1:
                    break
        if owner is not None and not owner["linkedin"]:
            owner["linkedin"] = href
    return people


def match_person_email(name: str, emails: list[str]) -> str:
    toks = [re.sub(r"[^a-z]", "", t.lower()) for t in name.split()]
    toks = [t for t in toks if t]
    if len(toks) < 2:
        return ""
    f, l = toks[0], toks[-1]
    pats = {f, l, f + l, f + "." + l, f + "_" + l, f[0] + l, f + l[0], f + "-" + l, l + f, f[0] + "." + l}
    for e in emails:
        if e.split("@")[0] in pats:
            return e
    return ""


def buyer_rank(title: str, service: str) -> int:
    order = BUYER_ORDER.get(service, BUYER_ORDER["default"])
    for i, rx in enumerate(order):
        if re.search(rx, title, re.I):
            return i
    return 99


# --------------------------------------------------------------------------- #
#  CRAWLER
# --------------------------------------------------------------------------- #
class Crawler:
    def __init__(self, cfg: dict):
        self.ua = cfg["discovery"]["user_agent"]
        self.delay = cfg["discovery"]["request_delay_seconds"]
        self.s = requests.Session()
        self.s.headers.update({"User-Agent": self.ua, "Accept-Language": "en"})
        self.robots: dict[str, robotparser.RobotFileParser] = {}
        self.use_careers = cfg["signals"]["use_hiring_hints"]

    def allowed(self, url: str) -> bool:
        p = urlparse(url)
        base = f"{p.scheme}://{p.netloc}"
        if base not in self.robots:
            rp = robotparser.RobotFileParser()
            try:
                r = self.s.get(base + "/robots.txt", timeout=8)
                rp.parse(r.text.splitlines() if r.status_code == 200 else [])
            except requests.RequestException:
                rp.parse([])
            rp.modified()
            self.robots[base] = rp
        return self.robots[base].can_fetch(self.ua, url)

    def get(self, url: str):
        if not self.allowed(url):
            return None
        try:
            r = self.s.get(url, timeout=12)
            if r.status_code != 200 or "html" not in r.headers.get("content-type", "").lower():
                return None
            sleep_range(self.delay)
            return r.text[:1_500_000]
        except requests.RequestException:
            return None

    def scrape(self, url: str) -> dict | None:
        p = urlparse(url)
        home = f"{p.scheme}://{p.netloc}/"
        domain = root_domain(home)
        html = self.get(home)
        if html is None:
            return None
        soup = BeautifulSoup(html, "html.parser")
        info_links, career_links = [], []
        for a in soup.find_all("a", href=True):
            full = urljoin(home, a["href"]).split("#")[0]
            if not full.startswith("http") or root_domain(full) != domain or full == home:
                continue
            low = (a["href"] + " " + a.get_text(" ", strip=True)).lower()
            if any(h in low for h in CAREER_HINTS):
                if self.use_careers and full not in career_links:
                    career_links.append(full)
            elif any(h in low for h in CONTACT_HINTS) and full not in info_links:
                info_links.append(full)
        pages = [(home, html, "home")]
        for link, kind in [(l, "info") for l in info_links[:4]] + [(l, "career") for l in career_links[:2]]:
            h = self.get(link)
            if h:
                pages.append((link, h, kind))

        raw_emails, phones, linkedin, text_all = [], [], "", ""
        title = og_name = desc = ""
        jobs, people, emp_disp, emp_est = [], [], "", 0
        for i, (pg_url, h, kind) in enumerate(pages):
            s = BeautifulSoup(h, "html.parser")
            if i == 0:
                og = s.find("meta", property="og:site_name")
                og_name = og["content"].strip() if og and og.get("content") else ""
                title = s.title.get_text(strip=True) if s.title else ""
                md = s.find("meta", attrs={"name": "description"}) or s.find("meta", property="og:description")
                desc = md["content"].strip()[:300] if md and md.get("content") else ""
            for a in s.find_all("a", href=True):
                href = a["href"]
                if href.lower().startswith("mailto:"):
                    raw_emails.append(href[7:].split("?")[0])
                elif href.lower().startswith("tel:"):
                    phones.append(href[4:].strip())
                elif "linkedin.com/company" in href and not linkedin:
                    linkedin = href.split("?")[0]
            for el in s.find_all(attrs={"data-cfemail": True}):
                raw_emails.append(decode_cf_email(el["data-cfemail"]))
            for it in jsonld_items(s):
                ne = it.get("numberOfEmployees")
                if ne and not emp_est:
                    if isinstance(ne, dict):
                        lo, hi = _to_int(ne.get("minValue", 0)), _to_int(ne.get("maxValue", 0))
                        v = _to_int(ne.get("value", 0))
                        if lo and hi:
                            emp_disp, emp_est = f"{lo}-{hi}", (lo + hi) // 2
                        elif v:
                            emp_disp, emp_est = str(v), v
                    elif _to_int(ne):
                        emp_disp, emp_est = str(_to_int(ne)), _to_int(ne)
                if it.get("email"):
                    raw_emails.append(str(it["email"]).replace("mailto:", ""))
            jobs += extract_jobs(s, pg_url, text_ok=(kind == "career"))
            if kind != "career":
                for person in extract_people(h):
                    if person["name"].lower() not in {x["name"].lower() for x in people}:
                        people.append(person)
            text = s.get_text(" ", strip=True)
            text_all += " " + text
            deob = re.sub(r"\s*[\[\(]\s*at\s*[\]\)]\s*", "@", text, flags=re.I)
            deob = re.sub(r"\s*[\[\(]\s*dot\s*[\]\)]\s*", ".", deob, flags=re.I)
            raw_emails += EMAIL_RE.findall(deob)
        if not phones:
            m = PHONE_RE.search(text_all)
            if m:
                phones.append(m.group(0).strip())
        if not emp_est:
            emp_disp, emp_est = parse_employees(text_all)

        name = og_name
        if not name and title:
            name = re.split(r"\s[|\-\u2013\u2014:]\s", title)[0].strip()
        if not name or len(name) > 60:
            name = domain.split(".")[0].replace("-", " ").title()
        return {"domain": domain, "website": home, "name": name, "emails": clean_emails(raw_emails, domain),
                "service_emails": clean_emails(raw_emails, domain, True), "phone": phones[0] if phones else "", "linkedin": linkedin, "description": desc,
                "text": text_all.lower()[:60000], "people": people, "jobs": jobs,
                "employees": emp_disp, "employees_est": emp_est}


# --------------------------------------------------------------------------- #
#  SIGNALS: which service do they need, and why?
# --------------------------------------------------------------------------- #
def mentions(name: str, domain: str, text: str) -> bool:
    t = text.lower()
    n = LEGAL_SUFFIX.sub("", name).strip(" .,-").lower()
    if len(n) >= 4 and n in t:
        return True
    stem = domain.split(".")[0]
    return len(stem) >= 5 and stem in t.replace(" ", "").replace("-", "")


def job_service(title: str) -> str:
    if JOB_GRC_RE.search(title):
        return "grc"
    if JOB_VAPT_RE.search(title):
        return "vapt"
    if re.search(r"soc|incident|threat|analyst|monitor", title, re.I):
        return "soc"
    return "general"


def need_text(cfg: dict, services: list, evidence: str) -> str:
    labels = " + ".join(cfg["services"].get(x, cfg["services"]["general"])["label"] for x in services)
    return f"Needs {labels}: {evidence}"


def industry_profile(cfg: dict, industry: str) -> dict:
    needs, low = cfg["industry_needs"], (industry or "").lower()
    for key, prof in needs.items():
        if key != "default" and key in low:
            return prof
    return needs["default"]


def market_of(location: str, domain: str) -> str:
    national = (location or "").lower() in RUNTIME["national"] or (domain or "").endswith(".pk")
    return "National" if national else "International"


def website_signals(cfg: dict, info: dict) -> list[dict]:
    sigs = []
    if cfg["signals"]["use_hiring_hints"]:
        for j in info["jobs"]:
            svc = job_service(j["title"])
            sigs.append({"type": "hiring", "service": svc, "services": [svc], "detail": j["title"],
                         "text": need_text(cfg, [svc], f"investing in security (hiring {j['title']})"),
                         "date": j["date"], "url": j["url"], "strength": 5, "source": "careers page"})
    found = []
    for key, label in FRAMEWORKS:
        if key in info["text"] and label not in found:
            found.append(label)
    if found:
        sigs.append({"type": "website_compliance", "service": "grc", "services": ["grc", "audit"],
                     "detail": ", ".join(found),
                     "text": need_text(cfg, ["grc", "audit"], f"publicly references {', '.join(found)}, so compliance is a live priority"),
                     "date": "", "url": info["website"], "strength": 3, "source": "website"})
    hits = [k for k in NEED_KEYWORDS if k in info["text"]]
    if len(hits) >= 2:
        sigs.append({"type": "attack_surface", "service": "vapt", "services": ["vapt"], "detail": ", ".join(hits[:3]),
                     "text": need_text(cfg, ["vapt"], f"runs public web / API / portal services ({', '.join(hits[:3])})"),
                     "date": "", "url": info["website"], "strength": 1, "source": "website"})
    return sigs


def search_signals(cfg: dict, name: str, domain: str) -> tuple[list[dict], str, int]:
    """Dated news + (optional) job-board + headcount lookups. Returns (signals, employees_display, employees_est)."""
    sc = cfg["signals"]
    n = LEGAL_SUFFIX.sub("", name).strip(" .,-") or name
    sigs, emp_disp, emp_est = [], "", 0

    def news(query, classify):
        for r in ddg_search("news", query, 6, "y"):
            txt = f"{r.get('title', '')} {r.get('body', '')}"
            date = parse_date(r.get("date"))
            if not mentions(name, domain, txt) or too_old(date, sc["hint_max_age_days"]):
                continue
            hit = classify(txt)
            if hit:
                title = (r.get("title") or "").strip()[:110]
                sigs.append({**hit, "service": hit["services"][0], "detail": title, "date": date,
                             "url": r.get("url") or "", "source": r.get("source") or "news",
                             "text": need_text(cfg, hit["services"], f"{hit['label']}: {title}")})
        sleep_range(sc["search_delay_seconds"])

    if sc["search_news"]:
        news(f'"{n}" data breach OR cyber attack OR ransomware',
             lambda t: {"type": "incident", "services": ["soc", "vapt"], "strength": 5,
                        "label": "security incident reported"} if INCIDENT_RE.search(t) else None)
        news(f'"{n}" ISO 27001 OR PCI DSS OR compliance OR launches OR funding OR expansion',
             lambda t: {"type": "compliance", "services": ["grc", "audit"], "strength": 4,
                        "label": "compliance / certification in the news"} if COMPLIANCE_RE.search(t) else
             ({"type": "launch", "services": ["vapt"], "strength": 3, "label": "new launch / growth"}
              if LAUNCH_RE.search(t) else None))
    if sc["use_hiring_hints"]:
        for r in ddg_search("text", f'"{n}" hiring security analyst OR "security engineer" OR pentester OR "compliance officer" OR "SOC analyst"', 8):
            title = r.get("title") or ""
            m = SEC_JOB_RE.search(title)
            if m and mentions(name, domain, f"{title} {r.get('body', '')}"):
                role = m.group(0).title()
                svc = job_service(role)
                sigs.append({"type": "hiring", "service": svc, "services": [svc], "detail": role,
                             "text": need_text(cfg, [svc], f"investing in security (hiring {role})"), "date": "",
                             "url": r.get("href") or "", "strength": 4, "source": "job board"})
                break
        sleep_range(sc["search_delay_seconds"])
    if sc["search_employees"]:
        for r in ddg_search("text", f'"{n}" company size employees LinkedIn', 6):
            txt = f"{r.get('title', '')} {r.get('body', '')}"
            if mentions(name, domain, txt):
                disp, est = parse_employees(txt)
                if est:
                    emp_disp, emp_est = disp, est
                    break
        sleep_range(sc["search_delay_seconds"])
    return sigs, emp_disp, emp_est


def build_lead_fields(cfg: dict, info: dict, cand: dict, industry: str, location: str, use_search: bool = True) -> dict:
    """Turns a scraped website into all lead fields (contact, hint, service, employees, decision-maker, score)."""
    domain, name = info["domain"], cand.get("name") or info["name"]
    emails = info["emails"]
    company_email = next((e for pref in GENERIC_LOCALS for e in emails if e.split("@")[0] == pref), emails[0] if emails else "")

    sigs = website_signals(cfg, info)
    emp_disp, emp_est, emp_src = info["employees"], info["employees_est"], "website (self-reported)"
    sc = cfg["signals"]
    if sc["enabled"] and use_search and (emails or not sc["only_for_leads_with_email"]):
        s2, d2, e2 = search_signals(cfg, name, domain)
        sigs += s2
        if not emp_est and e2:
            emp_disp, emp_est, emp_src = d2, e2, "web search snippet (approx.)"
    if not emp_est:
        emp_disp, emp_src = "", ""
    prof = industry_profile(cfg, industry)  # always add: what companies in this industry typically need
    sigs.append({"type": "industry", "service": prof["services"][0], "services": prof["services"], "detail": industry,
                 "text": need_text(cfg, prof["services"], prof["reason"]), "date": "", "url": "", "strength": 2,
                 "source": "industry profile"})
    sigs.sort(key=lambda s: (s["strength"], 1 if s["date"] else 0, s["date"]), reverse=True)
    top = sigs[0]
    labels, seen = [], set()
    for sg in sigs:
        for svc in sg["services"]:
            if svc not in seen:
                seen.add(svc)
                labels.append(cfg["services"].get(svc, cfg["services"]["general"])["label"])

    # decision-maker: best-ranked leader listed on the site, plus their email if it is published
    for p_ in info["people"]:
        p_["email"] = match_person_email(p_["name"], emails)
    people = sorted(info["people"], key=lambda p: buyer_rank(p["title"], top["service"]) + (0 if p["email"] else 1.5))
    who = people[0] if people else {}
    contact_email = who.get("email", "")
    others = "; ".join(f"{p['name']} ({p['title']})" for p in people[1:6])
    send_to = contact_email or (emails[0] if emails else "")
    q = f'{name} CISO OR CTO OR "IT Manager" OR "Head of Security"'
    lead = {
        "company": name, "website": info["website"], "industry": industry, "location": location,
        "email": send_to, "market": market_of(location, domain), "all_emails": ", ".join(emails[:6]), "company_email": company_email,
        "phone": cand.get("phone") or info["phone"], "linkedin": info["linkedin"],
        "address": cand.get("address", ""), "description": info["description"],
        "need_service": top["service"], "need_services_all": ", ".join(labels), "hint_type": top["type"],
        "hint_detail": top["detail"], "hint_text": top["text"], "hint_date": top["date"], "hint_url": top["url"],
        "hint_found_on": datetime.now().date().isoformat(),
        "other_hints": "; ".join(s["text"] + (f" [{s['date']}]" if s["date"] else "") for s in sigs[1:5]),
        "signals_json": json.dumps(sigs, ensure_ascii=False),
        "employees": emp_disp, "employees_est": emp_est, "employees_source": emp_src,
        "contact_name": who.get("name", ""), "contact_title": who.get("title", ""), "contact_email": contact_email,
        "contact_linkedin": who.get("linkedin", ""), "other_contacts": others,
        **({"notes": "Skipped customer-service mailboxes: " + "; ".join(info["service_emails"][:4])}
           if info["service_emails"] else {}),
        "buyer_search_url": "https://www.linkedin.com/search/results/people/?keywords=" + quote_plus(q),
    }
    lead["score"] = score_lead(cfg, lead, top["strength"])
    return lead


def score_lead(cfg: dict, lead: dict, strength: int = 0) -> int:
    s = 0
    if lead.get("email"):
        s += 30
        dom = lead["email"].split("@")[1]
        d = root_domain(lead.get("website") or "")
        if d and (dom == d or dom.endswith("." + d)):
            s += 10
    if lead.get("phone"):
        s += 5
    if lead.get("contact_name"):
        s += 10
    if lead.get("employees_est"):
        s += 5
    s += min(strength * 6, 30)
    if any(k in (lead.get("industry") or "").lower() for k in cfg["discovery"]["priority_industries"]):
        s += 10
    return min(s, 100)


# --------------------------------------------------------------------------- #
#  DISCOVERY / ENRICHMENT
# --------------------------------------------------------------------------- #
def pick_queries(cfg: dict, db: DB):
    """Random mix of national and international searches across every industry (no fixed area)."""
    d = cfg["discovery"]
    cutoff = (datetime.now() - timedelta(days=d["requery_after_days"])).isoformat()
    nat, intl = d["national_locations"], d["international_locations"]
    out, tried = [], set()
    for _ in range(600):
        if len(out) >= d["max_queries_per_run"] or not (nat or intl):
            break
        pool = nat if nat and (not intl or random.random() < d["national_ratio"]) else intl
        ind, loc, tpl = random.choice(d["industries"]), random.choice(pool), random.choice(d["query_templates"])
        q = (tpl if loc else "{industry} company contact us email").format(industry=ind, location=loc)
        q = re.sub(r"\s+", " ", q).strip()
        if q in tried:
            continue
        tried.add(q)
        row = db.one("SELECT run_at FROM seen_queries WHERE query=?", (q,))
        if row and row["run_at"] > cutoff:
            continue
        out.append((ind, loc, q))
    return out


def find_leads(cfg: dict, db: DB, limit: int | None = None) -> int:
    d = cfg["discovery"]
    limit = limit or d["max_new_leads_per_run"]
    queries = pick_queries(cfg, db)
    if not queries:
        log.info("No un-used search queries left. Add industries/locations in config.json.")
        return 0
    crawler = Crawler(cfg)
    st = {"urls": 0, "skipped": 0, "crawled": 0, "unreachable": 0, "saved": 0, "with_email": 0, "with_buyer": 0}
    try:
        for industry, location, q in queries:
            if st["with_email"] >= limit:
                break
            log.info("Searching: %s", q)
            cands = []
            if d["google_places_api_key"]:
                cands += search_places(q, d["google_places_api_key"], d["results_per_query"])
            cands += search_ddg(q, d["results_per_query"])
            db.run("INSERT OR REPLACE INTO seen_queries VALUES (?,?)", (q, now_iso()))
            st["urls"] += len(cands)
            for c in cands:
                if st["with_email"] >= limit:
                    break
                domain = root_domain(c["url"])
                if not domain or is_blocked(domain, cfg) or db.lead_exists(domain):
                    st["skipped"] += 1
                    continue
                info = crawler.scrape(c["url"])
                if info is None:
                    st["unreachable"] += 1
                    db.add_lead({"domain": domain, "website": c["url"], "company": c["name"] or domain,
                                 "industry": industry, "location": location, "status": "unreachable",
                                 "source": c["source"]})
                    continue
                st["crawled"] += 1
                if info["emails"] and db.is_suppressed(info["emails"][0]):
                    continue
                fields = build_lead_fields(cfg, info, c, industry, location)
                est, mn, mx = fields["employees_est"], d["min_employees"], d["max_employees"]
                out_of_size = est and ((mn and est < mn) or (mx and est > mx))
                fields.update(domain=domain, source=c["source"],
                              status="skipped_size" if out_of_size else ("new" if fields["email"] else "no_email"))
                db.add_lead(fields)
                st["saved"] += 1
                if fields["status"] == "new":
                    st["with_email"] += 1
                if fields["contact_name"]:
                    st["with_buyer"] += 1
                log.info("  + %s | %s | hint: %s%s | staff: %s | contact: %s | %s",
                         fields["company"], fields["need_services_all"], fields["hint_text"],
                         f" [{fields['hint_date']}]" if fields["hint_date"] else "",
                         fields["employees"] or "?", fields["contact_name"] or "-", fields["email"] or "no public email")
            export_leads(db, quiet=True)  # saved after every query, so nothing is lost if you stop the script
    finally:
        export_leads(db)
    log.info("Discovery summary: %d search results, %d already-known/blocked, %d sites crawled, %d unreachable, "
             "%d leads saved (%d with email, %d with a named decision-maker).",
             st["urls"], st["skipped"], st["crawled"], st["unreachable"], st["saved"], st["with_email"], st["with_buyer"])
    if st["urls"] == 0:
        log.error("The search engine returned 0 results for every query, so nothing could be saved. Likely causes: "
                  "no internet / VPN or firewall blocking, a temporary rate-limit (wait 15-30 min), or an outdated "
                  "package: pip install -U ddgs.  Last error: %s", SEARCH_STATS["last_error"] or "none")
        log.error("Run:  python rynex_leadgen.py doctor   for a full diagnosis.")
    return st["with_email"]


def enrich_existing(cfg: dict, db: DB, limit: int | None = None, force: bool = False) -> int:
    """Adds hint / service / employees / decision-maker to leads found before v2 (or re-does them with --force)."""
    sql = "SELECT * FROM leads WHERE status IN ('new','no_email','contacted','completed','replied')"
    if not force:
        sql += " AND (hint_type IS NULL OR hint_type='')"
    rows = db.all(sql + " ORDER BY id")
    if limit:
        rows = rows[:limit]
    crawler, done = Crawler(cfg), 0
    log.info("Enriching %d lead(s)...", len(rows))
    try:
        for r in rows:
            info = crawler.scrape(r["website"])
            if info is None:
                log.info("  ! %s unreachable, skipped", r["company"])
                continue
            f = build_lead_fields(cfg, info, {"name": r["company"], "phone": r["phone"], "address": r["address"]},
                                  r["industry"], r["location"])
            if r["status"] in ("new", "no_email"):
                f["status"] = "new" if f["email"] else "no_email"
            else:  # already contacted: never change who we email
                f.pop("email", None)
            f["company"] = r["company"]
            if r["notes"]:
                f.pop("notes", None)  # keep reply / follow-up notes
            db.update_lead(r["id"], **f)
            done += 1
            log.info("  + %s | %s | %s", r["company"], f["need_services_all"], f["hint_text"])
    finally:
        export_leads(db)
    return done


def import_csv(db: DB, path: str, cfg: dict) -> int:
    n = 0
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            row = {k.strip().lower(): (v or "").strip() for k, v in row.items() if k}
            email = row.get("email", "").lower()
            site = row.get("website") or row.get("domain") or (email.split("@")[1] if "@" in email else "")
            domain = root_domain(site)
            if not domain or (email and db.is_suppressed(email)):
                continue
            ind, loc = row.get("industry", ""), row.get("location", "")
            prof = industry_profile(cfg, ind)
            lead = {"domain": domain, "company": row.get("company") or domain, "website": site, "industry": ind,
                    "location": loc, "market": market_of(loc, domain), "email": email, "all_emails": email,
                    "company_email": email, "phone": row.get("phone", ""), "linkedin": row.get("linkedin", ""),
                    "source": "import", "status": "new" if email else "no_email", "need_service": prof["services"][0],
                    "need_services_all": ", ".join(cfg["services"][x]["label"] for x in prof["services"]),
                    "hint_type": "industry", "hint_detail": ind, "hint_text": need_text(cfg, prof["services"], prof["reason"]),
                    "hint_found_on": datetime.now().date().isoformat()}
            lead["score"] = score_lead(cfg, lead, 2)
            n += db.add_lead(lead)
    log.info("Imported %d leads from %s (run `enrich` to add decision-makers / news hints)", n, path)
    return n


# --------------------------------------------------------------------------- #
#  EMAIL: context-aware text + your HTML template
# --------------------------------------------------------------------------- #
def industry_pitch(cfg: dict, industry: str) -> str:
    return industry_profile(cfg, industry)["pitch"]


def pick_hook(cfg: dict, lead: dict) -> str:
    t, hooks = lead.get("hint_type") or "", cfg["hooks"]
    if t == "incident" and not cfg["signals"]["cite_incident_news"]:
        t = "industry"  # never mention a breach unless the owner explicitly enabled it
    if t in ("hiring", "incident", "compliance", "launch", "website_compliance") and lead.get("hint_detail"):
        return hooks[t]
    return hooks["industry"]


def template_vars(cfg: dict, lead: dict) -> SafeDict:
    c = cfg["company"]
    svc = cfg["services"].get(lead.get("need_service") or "general") or cfg["services"]["general"]
    company = lead.get("company") or lead.get("domain") or "your"
    contact, cemail = lead.get("contact_name") or "", (lead.get("contact_email") or "").lower()
    personal = bool(contact and cemail and (lead.get("email") or "").lower() == cemail)
    est = lead.get("employees_est") or 0
    size_key = "unknown" if not est else "small" if est < 50 else "medium" if est < 250 else "large"
    v = SafeDict(
        company=company, domain=lead.get("domain", ""), industry=lead.get("industry", ""),
        industry_pitch=industry_pitch(cfg, lead.get("industry", "")), hint_detail=lead.get("hint_detail") or "",
        sender_name=c["sender_name"], sender_title=c["sender_title"], sender_company=c["name"],
        website=c["website"], from_email=cfg["smtp"]["from_email"], postal_address=c["postal_address"],
        booking_line=f"You can pick a time here: {c['booking_link']}" if c.get("booking_link") else "",
        greeting=f"Hi {contact.split()[0]}," if personal else f"Hi {company} team,",
        forward_line="" if personal or not contact else
        f"I believe {contact} ({lead.get('contact_title') or 'leadership'}) may be the right person to speak to about this. If so, please pass it along.",
        size_line=cfg["size_lines"].get(size_key, ""), service_label=svc["label"], service_phrase=svc["phrase"],
        service_offer=svc["offer"], service_deliverable=svc["deliverable"], service_checklist=svc["checklist"])
    v["subject"] = svc["subject"].format_map(v)
    v["hook"] = pick_hook(cfg, lead).format_map(v)
    return v


def load_html_template(cfg: dict) -> str:
    path = BASE_DIR / cfg.get("email_template_file", "email_template.html")
    if not path.exists():
        path.write_text(DEFAULT_EMAIL_HTML, encoding="utf-8")
    return path.read_text(encoding="utf-8")


def body_to_html(body: str) -> str:
    out = []
    for i, para in enumerate(p for p in re.split(r"\n\s*\n", body.strip()) if p.strip()):
        txt = htmllib.escape(para.strip())
        txt = re.sub(r"(https?://[^\s<]+)", r'<a href="\1" style="color:rgb(37,99,235);text-decoration:none;">\1</a>', txt)
        color, size = ("rgb(31,41,55)", 18) if i == 0 else ("rgb(55,65,81)", 16)
        out.append(f'<p style="margin:0px 0px 16px;line-height:1.7;"><span style="color:{color};font-size:{size}px;">'
                   f'{txt.replace(chr(10), "<br>")}</span></p>')
    return "\n".join(out)


def render_email(cfg: dict, lead: dict, step: int):
    """Returns (subject, plain_text, html) for email number `step` (0 = intro), written for this lead's context."""
    t, v = cfg["templates"], template_vars(cfg, lead)
    initial_subject = t["initial"]["subject"].format_map(v)
    tpl = t["initial"] if step == 0 else t["followups"][step - 1]
    subject = initial_subject if step == 0 else (tpl.get("subject", "").format_map(v) or "Re: " + initial_subject)
    body = re.sub(r"[ \t]+\n", "\n", tpl["body"].format_map(v))
    body = re.sub(r"\n{3,}", "\n\n", body).strip()
    footer = cfg["compliance"]["footer"].format_map(v)
    text = f"{body}\n\n--\n{footer}"
    footer_html = "".join(f"<div>{htmllib.escape(line)}</div>" for line in footer.splitlines())
    preheader = htmllib.escape(v["hook"][:110])
    html = (load_html_template(cfg).replace("{{PREHEADER}}", preheader).replace("{{BODY}}", body_to_html(body))
            .replace("{{FOOTER}}", footer_html).replace("{{CONTACT_EMAIL}}", cfg["smtp"]["from_email"])
            .replace("{{WEBSITE}}", cfg["company"]["website"]))
    return subject, text, html


def build_message(cfg: dict, lead: dict, subject, text, html, step: int) -> EmailMessage:
    s, c = cfg["smtp"], cfg["company"]
    msg = EmailMessage()
    msg["From"] = formataddr((f"{c['sender_name']} | {c['name']}", s["from_email"]))
    msg["To"] = lead["email"]
    msg["Subject"] = subject
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain=s["from_email"].split("@")[-1])
    if s.get("reply_to"):
        msg["Reply-To"] = s["reply_to"]
    msg["List-Unsubscribe"] = f"<mailto:{s['from_email']}?subject=unsubscribe>"
    if step > 0 and lead.get("last_message_id"):
        msg["In-Reply-To"] = lead["last_message_id"]
        msg["References"] = f"{lead.get('first_message_id') or ''} {lead['last_message_id']}".strip()
    msg.set_content(text)
    msg.add_alternative(html, subtype="html")
    return msg


def smtp_connect(cfg: dict):
    s = cfg["smtp"]
    if s["security"] != "none" and not (s["username"] and s["password"]):
        sys.exit("SMTP login details missing. Put your Gmail App Password in config.json -> smtp.password "
                 "(or set RYNEX_SMTP_PASSWORD and reopen CMD), and make sure smtp.username is set.")
    ctx = ssl.create_default_context()
    if s["security"] == "ssl":
        srv = smtplib.SMTP_SSL(s["host"], s["port"], timeout=60, context=ctx)
    else:
        srv = smtplib.SMTP(s["host"], s["port"], timeout=60)
        srv.ehlo()
        if s["security"] == "starttls":
            srv.starttls(context=ctx)
            srv.ehlo()
    if s["username"] and s["password"]:
        srv.login(s["username"], s["password"])
    return srv


def validate_for_sending(cfg: dict):
    problems = []
    c, s = cfg["company"], cfg["smtp"]
    if "YOUR" in c["sender_name"].upper():
        problems.append("company.sender_name is still the placeholder")
    if "YOUR" in c["postal_address"].upper():
        problems.append("company.postal_address is still the placeholder (a physical address is legally required in commercial email in many countries)")
    if s["security"] != "none" and not s["password"]:
        problems.append("smtp.password is empty (use an App Password or env var RYNEX_SMTP_PASSWORD)")
    if problems:
        sys.exit("Fix config.json before sending:\n  - " + "\n  - ".join(problems))


def in_send_window(cfg: dict, now: datetime) -> bool:
    w = cfg["sending"]["send_window"]
    return now.weekday() in w["days"] and w["start_hour"] <= now.hour < w["end_hour"]


def send_simple(cfg: dict, to: str, subject: str, body: str):
    msg = EmailMessage()
    msg["From"] = cfg["smtp"]["from_email"]
    msg["To"] = to
    msg["Subject"] = subject
    msg.set_content(body)
    srv = smtp_connect(cfg)
    srv.send_message(msg)
    srv.quit()


def send_due(cfg: dict, db: DB, dry_run: bool = False, limit: int | None = None) -> int:
    sc = cfg["sending"]
    now = datetime.now()
    if not dry_run:
        validate_for_sending(cfg)
        if not in_send_window(cfg, now):
            log.info("Outside send window (%s). Nothing sent.", sc["send_window"])
            return 0
    sent_today = db.one("SELECT COUNT(*) c FROM emails_sent WHERE sent_at >= ?", (now.strftime("%Y-%m-%d"),))["c"]
    remaining = sc["daily_limit"] - sent_today
    if limit:
        remaining = min(remaining, limit)
    if remaining <= 0:
        log.info("Daily limit reached (%d).", sc["daily_limit"])
        return 0

    due = db.all("SELECT * FROM leads WHERE status='contacted' AND next_send_at IS NOT NULL "
                 "AND next_send_at<=? AND fail_count<3 ORDER BY next_send_at", (now_iso(),))
    new = db.all("SELECT * FROM leads WHERE status='new' AND email IS NOT NULL AND email!='' AND score>=? "
                 "AND fail_count<3 ORDER BY score DESC, id", (sc["min_lead_score"],))
    queue = [dict(r) for r in list(due) + list(new) if not db.is_suppressed(r["email"])][:remaining]
    log.info("Queue: %d email(s) (%d follow-ups due, %d new leads waiting).", len(queue), len(due), len(new))

    srv, sent = None, 0
    fu_days, fu_tpls = sc["followup_days"], cfg["templates"]["followups"]
    for i, lead in enumerate(queue):
        step = lead["step"]
        subject, text, html = render_email(cfg, lead, step)
        if dry_run:
            print(f"\n--- DRY RUN: email #{step + 1} to {lead['email']} ({lead['company']}) "
                  f"[service: {lead.get('need_service')}, hint: {lead.get('hint_type')}] ---\n"
                  f"Subject: {subject}\n\n{text}\n")
            continue
        msg = build_message(cfg, lead, subject, text, html, step)
        try:
            if srv is None:
                srv = smtp_connect(cfg)
            srv.send_message(msg)
        except smtplib.SMTPAuthenticationError as e:
            log.error("SMTP login failed: %s. Check username / App Password.", e)
            break
        except smtplib.SMTPRecipientsRefused:
            log.warning("Recipient refused, marking bounced: %s", lead["email"])
            db.update_lead(lead["id"], status="bounced", next_send_at=None)
            db.suppress(lead["email"], "bounced")
            continue
        except (smtplib.SMTPException, OSError) as e:
            log.warning("Send failed for %s: %s", lead["email"], e)
            db.update_lead(lead["id"], fail_count=lead["fail_count"] + 1)
            srv = None
            continue

        mid = msg["Message-ID"]
        sent_at = now_iso()
        db.run("INSERT INTO emails_sent (lead_id, step, to_email, subject, message_id, sent_at) VALUES (?,?,?,?,?,?)",
               (lead["id"], step, lead["email"], subject, mid, sent_at))
        fields = {"status": "contacted", "step": step + 1, "last_sent_at": sent_at,
                  "last_message_id": mid, "fail_count": 0}
        if step == 0:
            fields["first_message_id"] = mid
        if step < len(fu_days) and step < len(fu_tpls):
            fields["next_send_at"] = (datetime.now() + timedelta(days=fu_days[step])).isoformat(timespec="seconds")
        else:
            fields.update(status="completed", next_send_at=None)
        db.update_lead(lead["id"], **fields)
        sent += 1
        log.info("Sent email #%d to %s (%s, %s)", step + 1, lead["email"], lead["company"], lead.get("need_service"))
        if i < len(queue) - 1:
            sleep_range(sc["delay_between_emails_seconds"])
    if srv:
        try:
            srv.quit()
        except Exception:
            pass
    log.info("Sent %d email(s) this run.", sent)
    if sent:
        export_leads(db, quiet=True)
    return sent


# --------------------------------------------------------------------------- #
#  REPLY / BOUNCE / UNSUBSCRIBE DETECTION (IMAP)
# --------------------------------------------------------------------------- #
def get_text(msg) -> str:
    parts = msg.walk() if msg.is_multipart() else [msg]
    for p in parts:
        if p.get_content_type() == "text/plain":
            payload = p.get_payload(decode=True)
            if payload:
                return payload.decode(p.get_content_charset() or "utf-8", errors="replace")
    return ""


def fresh_reply_text(body: str) -> str:
    """Only the newly written part of a reply (drop quoted history, which contains OUR footer)."""
    out = []
    for line in body.splitlines():
        if line.strip().startswith(">") or re.match(r"^\s*(On .+wrote:|-{2,}\s*Original|From:\s|Sent from)", line, re.I):
            break
        out.append(line)
    return " ".join(out)[:600].lower()


def process_message(cfg: dict, db: DB, msg) -> str | None:
    own = cfg["smtp"]["from_email"].lower()
    from_addr = parseaddr(msg.get("From", ""))[1].lower()
    if not from_addr or from_addr == own or "@" not in from_addr:
        return None
    subject = decode_hdr(msg.get("Subject"))
    local, dom = from_addr.split("@", 1)

    if local in ("mailer-daemon", "postmaster") or re.search(
            r"undeliver|delivery status|failure notice|returned mail|delivery has failed", subject, re.I):
        hit = None
        for addr in set(EMAIL_RE.findall(msg.as_string())):
            lead = db.one("SELECT * FROM leads WHERE lower(email)=? AND status IN ('contacted','new')", (addr.lower(),))
            if lead:
                db.update_lead(lead["id"], status="bounced", next_send_at=None)
                db.suppress(addr, "bounced")
                log.info("Bounce detected: %s", addr)
                hit = "bounce"
        return hit

    if msg.get("Auto-Submitted", "no").lower() != "no" or re.search(
            r"auto(matic)?[\s\-]?reply|out of office|autoreply|automatic response", subject, re.I):
        return None

    lead = db.one("SELECT * FROM leads WHERE lower(email)=?", (from_addr,))
    if not lead:
        for ref in (msg.get("In-Reply-To", "") + " " + msg.get("References", "")).split():
            row = db.one("SELECT lead_id FROM emails_sent WHERE message_id=?", (ref,))
            if row:
                lead = db.one("SELECT * FROM leads WHERE id=?", (row["lead_id"],))
                break
    if not lead:
        lead = db.one("SELECT * FROM leads WHERE domain=? OR ? LIKE '%.' || domain", (dom, dom))
    if not lead or lead["status"] in ("replied", "unsubscribed", "bounced", "not_interested"):
        return None
    if lead["status"] == "new":
        return None  # we never contacted them

    reply = fresh_reply_text(get_text(msg))
    if any(p in reply for p in UNSUB_PHRASES):
        status = "unsubscribed"
    elif any(p in reply for p in DECLINE_PHRASES):
        status = "not_interested"
    else:
        status = "replied"
    db.update_lead(lead["id"], status=status, replied_at=now_iso(), next_send_at=None,
                   notes=((lead["notes"] or "") + f" [{now_iso()}] {status}: {reply[:200]}").strip())
    if status != "replied":
        db.suppress(from_addr, status)
        if lead["email"]:
            db.suppress(lead["email"], status)
    log.info("Reply from %s (%s) -> %s", from_addr, lead["company"], status.upper())

    notify = cfg["automation"].get("notify_email")
    if notify:
        try:
            send_simple(cfg, notify, f"[Rynex Lead] {status.upper()}: {lead['company']}",
                        f"Company: {lead['company']}\nWebsite: {lead['website']}\nFrom: {from_addr}\n"
                        f"Subject: {subject}\nPhone: {lead['phone']}\n\n{reply}")
        except Exception as e:
            log.warning("Could not send notification: %s", e)
    return status


def check_replies(cfg: dict, db: DB) -> int:
    ic = cfg["imap"]
    if not (ic["host"] and ic["username"] and ic["password"]):
        log.warning("IMAP not configured; skipping reply detection (follow-ups will not stop on replies!).")
        return 0
    events = 0
    try:
        M = imaplib.IMAP4_SSL(ic["host"], ic["port"])
        M.login(ic["username"], ic["password"])
        M.select(ic["folder"], readonly=True)
        since = (datetime.now() - timedelta(days=ic["lookback_days"])).strftime("%d-%b-%Y")
        typ, data = M.search(None, "SINCE", since)
        for num in data[0].split()[-400:]:
            typ, d = M.fetch(num, "(BODY.PEEK[])")
            if typ == "OK" and d and d[0]:
                if process_message(cfg, db, message_from_bytes(d[0][1])):
                    events += 1
        M.logout()
    except Exception as e:
        log.error("IMAP check failed: %s", e)
    log.info("Reply check finished: %d event(s).", events)
    return events


# --------------------------------------------------------------------------- #
#  EXPORT / STATS / CYCLE
# --------------------------------------------------------------------------- #
def export_rows(db: DB):
    rows = []
    for r in db.all("SELECT * FROM leads ORDER BY score DESC, id"):
        row = []
        for header, col in EXPORT_COLUMNS:
            v = r[col] if col in r.keys() else ""
            v = "" if v is None else v
            if col == "hint_date" and not v and r["hint_type"]:
                v = "Undated (no post date published)"
            elif col == "contact_name" and not v:
                v = "Not found on website"
            elif col == "employees" and not v:
                v = "Not found"
            elif col == "market" and not v:
                v = market_of(r["location"], r["domain"])
            row.append(v)
        rows.append(row)
    return [h for h, _ in EXPORT_COLUMNS], rows


def _write_csv(path: Path, header, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


LINK_COLS = {"Website", "Hint Source URL", "Decision-Maker LinkedIn", "Find Decision-Maker (LinkedIn search)",
             "Company LinkedIn"}
WRAP_COLS = {"Need Hint (why they need it)", "Service Needed", "Other Hints", "Other Leaders Found",
             "Other Emails Found", "Notes"}
COL_WIDTH = {"Company": 28, "Website": 30, "Market": 14, "Location": 16, "Industry": 20, "Service Needed": 22,
             "Need Hint (why they need it)": 62, "Hint Posted Date": 16, "Other Hints": 50, "Decision-Maker": 24,
             "Decision-Maker Title": 26, "Decision-Maker Email": 30, "Decision-Maker LinkedIn": 30,
             "Other Leaders Found": 40, "Company Email": 28, "Send-To Email": 30, "Other Emails Found": 34,
             "Notes": 40, "Find Decision-Maker (LinkedIn search)": 30, "Hint Source URL": 30, "Company LinkedIn": 30}


def write_xlsx(path: Path, header, rows):
    from collections import Counter
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(header)
    for r in rows:
        ws.append(r)
    for ci, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(ci)].width = COL_WIDTH.get(h, 16)
        ws.cell(row=1, column=ci).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for ri in range(2, len(rows) + 2):
            c = ws.cell(row=ri, column=ci)
            c.alignment = Alignment(vertical="top", wrap_text=(h in WRAP_COLS))
            if h in LINK_COLS and isinstance(c.value, str) and c.value.startswith("http") and len(c.value) < 2000:
                c.hyperlink = c.value
                c.font = Font(color="0563C1", underline="single")
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "B2"
    if rows:  # a real Excel "Table": banded rows + filter buttons + sortable
        tab = Table(displayName="LeadsTable", ref=f"A1:{get_column_letter(len(header))}{len(rows) + 1}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
    sm = wb.create_sheet("Summary")
    sm.append(["Metric", "Count"])
    sm.append(["Total leads", len(rows)])
    for col in ("Status", "Market"):
        idx = header.index(col)
        for k, v in sorted(Counter(r[idx] for r in rows).items(), key=lambda kv: -kv[1]):
            sm.append([f"{col}: {k}", v])
    sm.column_dimensions["A"].width = 34
    for c in sm[1]:
        c.font = Font(bold=True)
    wb.save(path)


def export_leads(db: DB, quiet: bool = False):
    header, rows = export_rows(db)
    try:
        import openpyxl  # noqa: F401
        have_xlsx = True
    except ImportError:
        have_xlsx = False
        log.warning("openpyxl is not installed (pip install openpyxl), so a CSV is saved instead of an Excel table.")
    saved = None
    stamp = f"{datetime.now():%Y%m%d_%H%M%S}"
    if have_xlsx:
        saved = XLSX_PATH
        try:
            write_xlsx(XLSX_PATH, header, rows)
        except PermissionError:  # open in Excel / locked by OneDrive
            saved = DATA_DIR / f"leads_{stamp}.xlsx"
            write_xlsx(saved, header, rows)
            log.warning("leads.xlsx is open in Excel. Saved to %s instead - close Excel to keep one file.", saved.name)
    if RUNTIME["csv"] or not have_xlsx:
        try:
            _write_csv(CSV_PATH, header, rows)
            saved = saved or CSV_PATH
        except PermissionError:
            alt = DATA_DIR / f"leads_{stamp}.csv"
            _write_csv(alt, header, rows)
            saved = saved or alt
    if not quiet:
        log.info("Saved %d leads -> %s", len(rows), saved.resolve())


def print_stats(db: DB):
    print("\n=== Rynex lead pipeline ===")
    for r in db.all("SELECT status, COUNT(*) c FROM leads GROUP BY status ORDER BY c DESC"):
        print(f"  {r['status']:<14}{r['c']}")
    total = db.one("SELECT COUNT(*) c FROM emails_sent")["c"]
    today = db.one("SELECT COUNT(*) c FROM emails_sent WHERE sent_at >= ?", (datetime.now().strftime("%Y-%m-%d"),))["c"]
    replied = db.one("SELECT COUNT(*) c FROM leads WHERE status IN ('replied','not_interested','unsubscribed')")["c"]
    contacted = db.one("SELECT COUNT(*) c FROM leads WHERE step>0")["c"]
    print(f"  emails sent   {total} total, {today} today")
    print(f"  reply rate    {replied}/{contacted}" + (f" ({100 * replied / contacted:.0f}%)" if contacted else ""))
    print()


def run_cycle(cfg: dict, db: DB, dry_run: bool = False):
    log.info("=== Cycle start ===")
    check_replies(cfg, db)  # first: stop follow-ups for anyone who answered
    backlog = db.one("SELECT COUNT(*) c FROM leads WHERE status='new'")["c"]
    if backlog < cfg["discovery"]["pause_discovery_if_backlog_over"]:
        find_leads(cfg, db)
    else:
        log.info("Backlog of %d unsent leads, skipping discovery.", backlog)
    if not dry_run:
        check_replies(cfg, db)  # again, discovery can take a while
    send_due(cfg, db, dry_run=dry_run)
    export_leads(db)
    print_stats(db)
    log.info("=== Cycle end ===")


def daemon(dry_run: bool):
    setup_logging()
    while True:
        try:
            cfg = load_config()  # re-read so config edits apply without restart
            db = DB()
            run_cycle(cfg, db, dry_run)
            hrs = cfg["automation"]["cycle_hours"]
        except SystemExit:
            raise
        except Exception:
            log.exception("Cycle crashed; will retry in 1 hour.")
            hrs = 1
        log.info("Sleeping %.1f hours. Press Ctrl+C to stop.", hrs)
        time.sleep(hrs * 3600)


# --------------------------------------------------------------------------- #
#  CLI
# --------------------------------------------------------------------------- #
def cmd_init():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if CONFIG_PATH.exists():
        print(f"{CONFIG_PATH} already exists (not overwritten).")
    else:
        CONFIG_PATH.write_text(json.dumps(DEFAULT_CONFIG, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"Created {CONFIG_PATH}")
    tpl = BASE_DIR / DEFAULT_CONFIG["email_template_file"]
    if not tpl.exists():
        tpl.write_text(DEFAULT_EMAIL_HTML, encoding="utf-8")
        print(f"Created {tpl} (your email design - edit freely)")
    print("\nNext steps:\n"
          "  1. Edit config.json: company.sender_name, company.postal_address, smtp.password (App Password)\n"
          "  2. python rynex_leadgen.py doctor          (checks search, saving, email settings)\n"
          "  3. python rynex_leadgen.py find\n"
          "  4. python rynex_leadgen.py preview         (see the email + open data\\preview_email.html)\n"
          "  5. python rynex_leadgen.py send --dry-run, then: python rynex_leadgen.py daemon")


def cmd_doctor(cfg: dict):
    print(f"Project folder : {BASE_DIR}")
    print(f"Leads (Excel)  : {XLSX_PATH}")
    try:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        t = DATA_DIR / "_write_test.tmp"
        t.write_text("ok")
        t.unlink()
        print("Can write files: YES")
    except Exception as e:
        print(f"Can write files: NO ({e})  <- fix folder permissions / move out of OneDrive")
    print("\nTesting web search (this needs internet)...")
    res = ddg_search("text", "cybersecurity company Karachi", 5)
    if res:
        print(f"  Search OK: {len(res)} results, e.g. {res[0].get('href') or res[0].get('url')}")
    else:
        print(f"  Search FAILED. Last error: {SEARCH_STATS['last_error'] or 'no results returned'}")
        print("  Try: pip install -U ddgs   |  turn off VPN/proxy  |  wait 15-30 min if rate-limited")
    news = ddg_search("news", "cybersecurity Pakistan", 3, "m")
    print(f"  News search: {'OK' if news else 'no results (hints from news will be empty)'}")
    s, i = cfg["smtp"], cfg["imap"]
    print(f"\nSMTP: {s['host']}:{s['port']} user={s['username'] or '(none)'} password={'set' if s['password'] else 'MISSING'}")
    print(f"IMAP: {i['host']} user={i['username'] or '(none)'} password={'set' if i['password'] else 'MISSING'}")
    print(f"Sender name / address set: {'YES' if 'YOUR' not in cfg['company']['sender_name'].upper() and 'YOUR' not in cfg['company']['postal_address'].upper() else 'NO (placeholders still in config.json)'}")
    db = DB()
    print(f"\nDatabase: {DB_PATH}  ({db.one('SELECT COUNT(*) c FROM leads')['c']} leads)")


def cmd_preview(cfg: dict, db: DB, lead_id: int | None):
    row = db.one("SELECT * FROM leads WHERE id=?", (lead_id,)) if lead_id else \
        db.one("SELECT * FROM leads WHERE email IS NOT NULL AND email!='' ORDER BY score DESC LIMIT 1")
    if not row:
        sys.exit("No lead with an email yet. Run `find` or `import` first.")
    lead = dict(row)
    print(f"Lead: {lead['company']} | service: {lead.get('need_service')} | hint: {lead.get('hint_text')} | "
          f"employees: {lead.get('employees') or '?'} | contact: {lead.get('contact_name') or '-'}\n")
    for step in range(1 + len(cfg["templates"]["followups"])):
        subject, text, html = render_email(cfg, lead, step)
        print(f"===== EMAIL #{step + 1} =====\nSubject: {subject}\n\n{text}\n")
        if step == 0:
            out = DATA_DIR / "preview_email.html"
            out.write_text(html, encoding="utf-8")
            print(f"(HTML design saved to {out} - double-click to open in your browser)\n")


def main():
    ap = argparse.ArgumentParser(description="Rynex Security - automated lead generation & follow-ups")
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("init", help="create config.json + email_template.html")
    sub.add_parser("doctor", help="diagnose search / saving / email settings")
    p = sub.add_parser("find", help="discover new leads and save them")
    p.add_argument("--limit", type=int)
    p = sub.add_parser("enrich", help="add hints/employees/decision-makers to leads found earlier")
    p.add_argument("--limit", type=int)
    p.add_argument("--force", action="store_true", help="redo leads that already have hints")
    p = sub.add_parser("preview", help="show the emails a lead would receive")
    p.add_argument("--id", type=int)
    p = sub.add_parser("send", help="send due intro emails + follow-ups")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--limit", type=int)
    sub.add_parser("replies", help="check inbox for replies / bounces / unsubscribes")
    p = sub.add_parser("run", help="one full cycle (replies -> find -> send -> export)")
    p.add_argument("--dry-run", action="store_true")
    p = sub.add_parser("daemon", help="run cycles forever")
    p.add_argument("--dry-run", action="store_true")
    p = sub.add_parser("export", help="write the Excel table (add --csv for a CSV too)")
    p.add_argument("--csv", action="store_true")
    sub.add_parser("stats", help="show pipeline numbers")
    p = sub.add_parser("import", help="import leads from a CSV (company,website,email,phone,industry,location)")
    p.add_argument("file")
    p = sub.add_parser("unsubscribe", help="add an address to the do-not-contact list")
    p.add_argument("email")
    p = sub.add_parser("test-email", help="send the intro email to yourself to check delivery/format")
    p.add_argument("to")
    p.add_argument("--service", default="soc", choices=["vapt", "soc", "grc", "general"])
    a = ap.parse_args()

    if a.cmd == "init":
        return cmd_init()
    setup_logging()
    if a.cmd == "daemon":
        return daemon(a.dry_run)
    cfg, db = load_config(), DB()
    if a.cmd == "doctor":
        cmd_doctor(cfg)
    elif a.cmd == "find":
        find_leads(cfg, db, a.limit)
    elif a.cmd == "enrich":
        enrich_existing(cfg, db, a.limit, a.force)
    elif a.cmd == "preview":
        cmd_preview(cfg, db, a.id)
    elif a.cmd == "send":
        send_due(cfg, db, a.dry_run, a.limit)
    elif a.cmd == "replies":
        check_replies(cfg, db)
    elif a.cmd == "run":
        run_cycle(cfg, db, a.dry_run)
    elif a.cmd == "export":
        RUNTIME["csv"] = RUNTIME["csv"] or a.csv
        export_leads(db)
    elif a.cmd == "stats":
        print_stats(db)
    elif a.cmd == "import":
        import_csv(db, a.file, cfg)
        export_leads(db)
    elif a.cmd == "unsubscribe":
        db.suppress(a.email, "manual")
        db.run("UPDATE leads SET status='unsubscribed', next_send_at=NULL WHERE lower(email)=?", (a.email.lower(),))
        print(f"{a.email} will not be contacted again.")
    elif a.cmd == "test-email":
        fake = {"company": "Test Company", "domain": "example.com", "industry": "fintech", "email": a.to,
                "step": 0, "need_service": a.service, "hint_type": "hiring",
                "hint_detail": "Security Analyst", "employees_est": 40}
        subject, text, html = render_email(cfg, fake, 0)
        msg = build_message(cfg, fake, "[TEST] " + subject, text, html, 0)
        srv = smtp_connect(cfg)
        srv.send_message(msg)
        srv.quit()
        print(f"Test email sent to {a.to}. Tip: also send one to a mail-tester.com address to check spam score.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nStopped.")
